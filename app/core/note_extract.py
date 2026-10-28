"""Ekstraksi teks dari dokumen catatan kasar untuk sumber notula."""
from pathlib import Path

from app.core import pdf_extract


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class NoteExtractError(RuntimeError):
    """Dokumen catatan tidak dapat dibaca atau tidak berisi teks."""


def extract_note_text(path: str) -> str:
    """Ambil teks dari PDF, DOCX, TXT, atau Markdown dan rapikan hasilnya."""
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        try:
            return pdf_extract.extract_pdf_text(path)
        except pdf_extract.PDFError as e:
            raise NoteExtractError(str(e)) from e
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in {".txt", ".md"}:
        return _extract_plain_text(path)
    raise NoteExtractError("Format catatan tidak didukung.")


def _extract_docx(path: str) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise NoteExtractError(
            "Paket 'python-docx' belum terinstall. Jalankan: pip install python-docx"
        ) from e

    try:
        doc = Document(path)
    except Exception as e:
        raise NoteExtractError(f"Dokumen Word tidak dapat dibaca: {e}") from e

    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            # python-docx mengekspos sel hasil merge horizontal lebih dari sekali.
            # Gunakan elemen XML sebagai identitas agar teksnya tidak berulang.
            seen_cells = set()
            values = []
            for cell in row.cells:
                cell_id = id(cell._tc)
                if cell_id in seen_cells:
                    continue
                seen_cells.add(cell_id)
                if cell.text and cell.text.strip():
                    values.append(cell.text.strip())
            if values:
                parts.append(" | ".join(values))

    result = "\n\n".join(parts).strip()
    if not result:
        raise NoteExtractError("Dokumen Word tidak berisi teks yang dapat dibaca.")
    return result


def _extract_plain_text(path: str) -> str:
    raw = None
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            raw = Path(path).read_text(encoding=encoding)
            break
        except UnicodeError:
            continue
        except OSError as e:
            raise NoteExtractError(f"File catatan tidak dapat dibaca: {e}") from e

    result = (raw or "").strip()
    if not result:
        raise NoteExtractError("File catatan tidak berisi teks yang dapat dibaca.")
    return result
