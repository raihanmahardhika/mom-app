"""
Membangun file .docx MoM dengan format yang meniru template MoM RSM Indonesia.

Struktur:
  0. Header dokumen : logo RSM (2,69 x 1,13 cm) rata kanan — muncul di tiap halaman
  1. Tabel Header  : Risalah Rapat | Klien | Penugasan | Topik | Hari,Tanggal |
                     Jam | Metode Rapat | Lokasi | Kehadiran (2 grup)
  2. Blok Diskusi  : "Hasil diskusi:" + poin-poin bernomor & bullet
  3. Tabel Tindak Lanjut : No | Tindak Lanjut | Ref | PIC | Target
  4. Tabel Expected Next Meeting(s), bila rencana rapat berikutnya diisi

Semua tabel memakai lebar total yang sama (selebar area cetak) dengan
layout FIXED, sehingga tepinya lurus satu sama lain.
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- konfigurasi tampilan ----------
FONT_NAME = "Arial"          # semua teks memakai Arial
SIZE_TITLE = 18              # judul bar: "Risalah Rapat", "Diskusi", "Tindak Lanjut"
SIZE_BODY = 11               # SEMUA teks lainnya
DISCUSSION_LINE_SPACING = 1.5
LOGO_PATH = os.path.join(os.path.dirname(__file__), "..", "assets", "logo_rsm.png")
LOGO_W_CM = 2.69
LOGO_H_CM = 1.13
BULLET_STYLE = "List Bullet"   # style bullet asli bawaan Word

# Proporsi lebar kolom tiap tabel (dinormalisasi ke lebar area cetak).
BOBOT_HEADER = [1700, 5600, 2060]
BOBOT_TINDAK = [700, 4900, 1000, 1560, 1200]
BOBOT_NEXT = [700, 4000, 2000, 2660]

EMU_PER_TWIP = 635


def _set_run_font(run, size, bold=False, color=None):
    """Terapkan Arial (termasuk untuk East Asian/complex script) + ukuran."""
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FONT_NAME
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT_NAME)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


# ---------- helper low-level ----------

def _set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def _set_table_borders(table, sz=4, color="000000"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)


def _lebar_cetak(doc) -> int:
    """Lebar area cetak dalam twips (dxa) — dipakai semua tabel agar seragam."""
    sec = doc.sections[0]
    return int((sec.page_width - sec.left_margin - sec.right_margin) / EMU_PER_TWIP)


def _set_table_layout(table, total, bobot):
    """
    Kunci lebar tabel: layout FIXED + lebar total eksplisit + lebar tiap kolom.

    Tanpa ini Word memakai autofit dan menghitung ulang lebar kolom mengikuti isi,
    sehingga tabel-tabel jadi tidak sama lebar.

    Harus dipanggil SEBELUM sel di-merge.
    """
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    for tag in ("w:tblLayout", "w:tblW"):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(total))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)

    jml = sum(bobot)
    lebar = [int(total * b / jml) for b in bobot]
    lebar[-1] = total - sum(lebar[:-1])       # sisa pembulatan ke kolom terakhir

    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc, w in zip(grid.findall(qn("w:gridCol")), lebar):
            gc.set(qn("w:w"), str(w))
    for row in table.rows:
        for cell, w in zip(row.cells, lebar):
            cell.width = Twips(w)
    return lebar


def _rapatkan(p, before=0, after=0, line_spacing=1.0):
    """Buang jarak antar-paragraf bawaan Word (docDefaults: after 8pt, line 1.08)."""
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing
    return p


def _kosongkan_sel(cell):
    """
    Kosongkan sel dan sisakan tepat satu paragraf tanpa run.

    Dipakai sebagai ganti `cell.text = ""`, yang meninggalkan run kosong tanpa
    format — kalau user mengedit sel di Word, kursor bisa mendarat di run itu
    dan teks baru keluar sebagai Calibri, bukan Arial.
    """
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    return p


def _cell_text(cell, text, bold=False, size=SIZE_BODY, align=None, color=None,
               valign=WD_ALIGN_VERTICAL.CENTER):
    _kosongkan_sel(cell)
    if valign is not None:
        cell.vertical_alignment = valign
    p = cell.paragraphs[0]
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    _rapatkan(p)
    run = p.add_run(text if text is not None else "")
    _set_run_font(run, size, bold=bold, color=color)
    return p


def _add_para(cell_or_doc, text, bold=False, size=SIZE_BODY, bullet=False, num=None,
              line_spacing=1.0):
    """
    Paragraf isi diskusi.

    bullet=True memakai style "List Bullet" bawaan Word — bullet sungguhan yang
    terhubung ke definisi numbering, bukan karakter yang diketik manual.
    """
    if bullet:
        try:
            p = cell_or_doc.add_paragraph(style=BULLET_STYLE)
        except KeyError:
            p = cell_or_doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
    else:
        p = cell_or_doc.add_paragraph()

    # Rata kanan-kiri agar blok diskusi terlihat rapi saat teksnya panjang.
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _rapatkan(p, before=2, after=2, line_spacing=line_spacing)

    if num is not None:
        run = p.add_run(f"{num}. ")
        _set_run_font(run, size, bold=True)
    run = p.add_run(text)
    _set_run_font(run, size, bold=bold)
    return p


# ---------- builder utama ----------

def build_mom(data: dict, out_path: str):
    """
    data = {
      "klien": str,
      "penugasan": str,
      "topik": str,
      "hari_tanggal": str,
      "jam": str,
      "metode": str,
      "lokasi": str,
      "kehadiran": [ {"grup": str, "orang": [str, ...]}, ... ],
      "diskusi": [ {"judul": str, "narasi": [str, ...], "poin": [str, ...]}, ... ],
      "tindak_lanjut": [ {"tindak_lanjut": str, "ref": str, "pic": str, "target": str}, ... ],
      "next_meetings": [ {"nama": str, "tanggal": str, "catatan": str}, ... ],
    }
    """
    doc = Document()

    # style default seluruh dokumen -> Arial 11
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(SIZE_BODY)
    # Word default: space_after 8pt & line 1.08 -> bikin baris tabel terlihat renggang.
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT_NAME)

    # margin dokumen
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    total = _lebar_cetak(doc)

    _add_logo_header(doc)
    _build_header_table(doc, data, total)
    doc.add_paragraph()
    _build_diskusi(doc, data, total)
    doc.add_paragraph()
    _build_tindak_lanjut(doc, data, total)
    if _has_next_meetings(data.get("next_meetings", [])):
        doc.add_paragraph()
        _build_next_meetings(doc, data, total)

    doc.save(out_path)
    return out_path


def _add_logo_header(doc):
    """Logo RSM di HEADER dokumen, rata kanan, ukuran tetap 2,69 x 1,13 cm."""
    header = doc.sections[0].header
    header.is_linked_to_previous = False
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    if os.path.exists(LOGO_PATH):
        run = p.add_run()
        run.add_picture(LOGO_PATH, width=Cm(LOGO_W_CM), height=Cm(LOGO_H_CM))
    return p


def _build_header_table(doc, data, total):
    labels = [
        ("Klien", data.get("klien", "")),
        ("Penugasan", data.get("penugasan", "")),
        ("Topik", data.get("topik", "")),
        ("Hari, Tanggal", data.get("hari_tanggal", "")),
        ("Jam", data.get("jam", "")),
        ("Metode Rapat", data.get("metode", "")),
        ("Lokasi", data.get("lokasi", "")),
    ]
    kehadiran = data.get("kehadiran", [])

    # 1 baris judul + N baris field + baris kehadiran per grup
    total_rows = 1 + len(labels) + max(1, len(kehadiran))
    table = doc.add_table(rows=total_rows, cols=3)
    _set_table_borders(table)
    _set_table_layout(table, total, BOBOT_HEADER)

    # Baris 0: "Risalah Rapat" — merge 3 kolom, rata tengah, 18pt
    r0 = table.rows[0]
    merged = r0.cells[0].merge(r0.cells[1]).merge(r0.cells[2])
    _cell_text(merged, "Risalah Rapat", bold=True, size=SIZE_TITLE,
               align=WD_ALIGN_PARAGRAPH.CENTER, valign=WD_ALIGN_VERTICAL.CENTER)
    _set_cell_bg(merged, "D9E2F3")
    _set_cell_margins(merged, top=100, bottom=100)

    # Baris field — label & isi rata kiri, tengah secara vertikal
    ridx = 1
    for label, value in labels:
        row = table.rows[ridx]
        c_lbl = row.cells[0]
        c_val = row.cells[1].merge(row.cells[2])
        _cell_text(c_lbl, label, bold=True, size=SIZE_BODY,
                   align=WD_ALIGN_PARAGRAPH.LEFT, valign=WD_ALIGN_VERTICAL.CENTER)
        _set_cell_bg(c_lbl, "F2F2F2")
        _cell_text(c_val, value, size=SIZE_BODY,
                   align=WD_ALIGN_PARAGRAPH.LEFT, valign=WD_ALIGN_VERTICAL.CENTER)
        _set_cell_margins(c_lbl); _set_cell_margins(c_val)
        ridx += 1

    # Baris kehadiran — kolom nama & instansi diisi per grup
    if not kehadiran:
        kehadiran = [{"grup": "", "orang": []}]
    awal = ridx
    for grup in kehadiran:
        row = table.rows[ridx]
        c_names = row.cells[1]
        c_grup = row.cells[2]

        # daftar nama bernomor
        _kosongkan_sel(c_names)
        c_names.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        first = True
        for n_i, nama in enumerate(grup.get("orang", []), start=1):
            p = c_names.paragraphs[0] if first else c_names.add_paragraph()
            first = False
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _rapatkan(p)
            run = p.add_run(f"{n_i}. {nama}")
            _set_run_font(run, SIZE_BODY)

        _cell_text(c_grup, grup.get("grup", ""), size=SIZE_BODY,
                   align=WD_ALIGN_PARAGRAPH.LEFT, valign=WD_ALIGN_VERTICAL.CENTER)
        _set_cell_margins(c_names); _set_cell_margins(c_grup)
        ridx += 1
    akhir = ridx - 1

    # Label "Kehadiran" di-merge vertikal sepanjang jumlah grup instansi.
    # Teks ditulis SETELAH merge, karena merge menggabungkan isi tiap sel.
    sel_label = table.rows[awal].cells[0]
    for r in range(awal + 1, akhir + 1):
        sel_label = sel_label.merge(table.rows[r].cells[0])
    _cell_text(sel_label, "Kehadiran", bold=True, size=SIZE_BODY,
               align=WD_ALIGN_PARAGRAPH.LEFT, valign=WD_ALIGN_VERTICAL.CENTER)
    _set_cell_bg(sel_label, "F2F2F2")
    _set_cell_margins(sel_label)


def _build_diskusi(doc, data, total):
    """
    Judul "Diskusi" dan isinya berada dalam SATU tabel (2 baris, 1 kolom),
    sehingga lebarnya dijamin identik — bukan dua tabel terpisah yang
    kebetulan diberi lebar sama.
    """
    table = doc.add_table(rows=2, cols=1)
    _set_table_borders(table)
    _set_table_layout(table, total, [1])

    # baris 0: judul — rata tengah, 18pt
    c = table.rows[0].cells[0]
    _cell_text(c, "Diskusi", bold=True, size=SIZE_TITLE,
               align=WD_ALIGN_PARAGRAPH.CENTER, valign=WD_ALIGN_VERTICAL.CENTER)
    _set_cell_bg(c, "D9E2F3")
    _set_cell_margins(c, top=100, bottom=100)

    # baris 1: isi — rata kiri atas
    cell = table.rows[1].cells[0]
    _kosongkan_sel(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    _set_cell_margins(cell, top=120, bottom=120, left=150, right=150)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _rapatkan(p, before=0, after=2)
    run = p.add_run("Hasil diskusi:")
    _set_run_font(run, SIZE_BODY, bold=True)

    for i, sec in enumerate(data.get("diskusi", []), start=1):
        _add_para(cell, sec.get("judul", ""), bold=True, size=SIZE_BODY, num=i)
        for narasi in sec.get("narasi", []):
            _add_para(
                cell,
                narasi,
                size=SIZE_BODY,
                line_spacing=DISCUSSION_LINE_SPACING,
            )
        for poin in sec.get("poin", []):
            _add_para(
                cell,
                poin,
                size=SIZE_BODY,
                bullet=True,
                line_spacing=DISCUSSION_LINE_SPACING,
            )

    endp = cell.add_paragraph()
    endp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _rapatkan(endp, before=6, after=0)
    r = endp.add_run("======= End of Minutes =======")
    _set_run_font(r, SIZE_BODY, bold=True)


def _build_tindak_lanjut(doc, data, total):
    tl = data.get("tindak_lanjut", [])
    rows = 2 + len(tl)  # judul + header + isi
    table = doc.add_table(rows=rows, cols=5)
    _set_table_borders(table)
    _set_table_layout(table, total, BOBOT_TINDAK)

    # judul merge
    r0 = table.rows[0]
    merged = r0.cells[0]
    for j in range(1, 5):
        merged = merged.merge(r0.cells[j])
    # Gaya judul disamakan dengan "Risalah Rapat" & "Diskusi": 18pt, rata tengah.
    _cell_text(merged, "Tindak Lanjut", bold=True, size=SIZE_TITLE,
               align=WD_ALIGN_PARAGRAPH.CENTER, valign=WD_ALIGN_VERTICAL.CENTER)
    _set_cell_bg(merged, "D9E2F3")
    _set_cell_margins(merged, top=100, bottom=100)

    # header kolom
    heads = ["No", "Tindak Lanjut", "Ref", "PIC", "Target"]
    for i, h in enumerate(heads):
        _cell_text(table.rows[1].cells[i], h, bold=True, size=SIZE_BODY,
                   align=WD_ALIGN_PARAGRAPH.CENTER, valign=WD_ALIGN_VERTICAL.CENTER)
        _set_cell_bg(table.rows[1].cells[i], "F2F2F2")

    for idx, item in enumerate(tl, start=1):
        row = table.rows[1 + idx]
        _cell_text(row.cells[0], str(idx), size=SIZE_BODY,
                   align=WD_ALIGN_PARAGRAPH.CENTER, valign=WD_ALIGN_VERTICAL.TOP)
        _cell_text(row.cells[1], item.get("tindak_lanjut", ""), size=SIZE_BODY,
                   valign=WD_ALIGN_VERTICAL.TOP)
        _cell_text(row.cells[2], item.get("ref", ""), size=SIZE_BODY,
                   align=WD_ALIGN_PARAGRAPH.CENTER, valign=WD_ALIGN_VERTICAL.TOP)
        _cell_text(row.cells[3], item.get("pic", ""), size=SIZE_BODY,
                   valign=WD_ALIGN_VERTICAL.TOP)
        _cell_text(row.cells[4], item.get("target", ""), size=SIZE_BODY,
                   valign=WD_ALIGN_VERTICAL.TOP)
        for c in row.cells:
            _set_cell_margins(c)


def _build_next_meetings(doc, data, total):
    nm = [
        meeting for meeting in data.get("next_meetings", [])
        if isinstance(meeting, dict) and any(
            str(meeting.get(key, "") or "").strip()
            for key in ("nama", "tanggal", "catatan")
        )
    ]
    if not nm:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _rapatkan(p, before=0, after=4)
    r = p.add_run("Expected Next Meeting(s)")
    _set_run_font(r, SIZE_BODY, bold=True)
    r.italic = True

    rows = 1 + len(nm)
    table = doc.add_table(rows=rows, cols=4)
    _set_table_borders(table)
    _set_table_layout(table, total, BOBOT_NEXT)

    heads = ["No.", "Nama Meeting", "Tanggal", "Catatan"]
    for i, h in enumerate(heads):
        _cell_text(table.rows[0].cells[i], h, bold=True, size=SIZE_BODY,
                   align=WD_ALIGN_PARAGRAPH.CENTER, valign=WD_ALIGN_VERTICAL.CENTER)
        _set_cell_bg(table.rows[0].cells[i], "F2F2F2")

    for idx, m in enumerate(nm, start=1):
        row = table.rows[idx]
        _cell_text(row.cells[0], f"{idx}.", size=SIZE_BODY,
                   align=WD_ALIGN_PARAGRAPH.CENTER, valign=WD_ALIGN_VERTICAL.TOP)
        _cell_text(row.cells[1], m.get("nama", ""), size=SIZE_BODY,
                   valign=WD_ALIGN_VERTICAL.TOP)
        _cell_text(row.cells[2], m.get("tanggal", ""), size=SIZE_BODY,
                   valign=WD_ALIGN_VERTICAL.TOP)
        _cell_text(row.cells[3], m.get("catatan", ""), size=SIZE_BODY,
                   valign=WD_ALIGN_VERTICAL.TOP)


def _has_next_meetings(meetings) -> bool:
    """True hanya bila ada minimal satu rencana rapat yang berisi data."""
    return any(
        isinstance(meeting, dict) and any(
            str(meeting.get(key, "") or "").strip()
            for key in ("nama", "tanggal", "catatan")
        )
        for meeting in (meetings or [])
    )


if __name__ == "__main__":
    # Uji cepat: render kerangka MoM kosong untuk memastikan builder berjalan.
    #   python -m app.core.docx_builder
    kosong = {
        "klien": "", "penugasan": "", "topik": "", "hari_tanggal": "",
        "jam": "", "metode": "", "lokasi": "",
        "kehadiran": [], "diskusi": [], "tindak_lanjut": [], "next_meetings": [],
    }
    out = os.path.join(os.path.dirname(__file__), "..", "..", "data", "outputs", "kerangka_mom.docx")
    build_mom(kosong, os.path.abspath(out))
    print("OK ->", os.path.abspath(out))
